"""S10 document and OCR tool boundary tests."""

import io
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from app.core.errors import BusinessError
from app.services import file_tools, file_upload
from app.storage import ocr


def _xlsx(rows: int = 1) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(["name", "qty"])
    for index in range(rows):
        sheet.append([f"item-{index}", index])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def _docx() -> bytes:
    document = Document()
    document.add_paragraph("hello document")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "item"
    table.cell(0, 1).text = "3"
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _pdf(pages: int = 1) -> bytes:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=72, height=72)
    stream = io.BytesIO()
    writer.write(stream)
    return stream.getvalue()


def _owned(monkeypatch: pytest.MonkeyPatch, name: str, content: bytes, *, owner: int = 7) -> None:
    def read(user_id: int, file_id: int, *, max_bytes: int | None = None) -> tuple[SimpleNamespace, bytes]:
        if user_id != owner or file_id != 12:
            raise BusinessError(403, "File access denied", 403)
        if max_bytes is not None and len(content) > max_bytes:
            raise BusinessError(400, "File exceeds read limit")
        return SimpleNamespace(original_name=name), content

    monkeypatch.setattr(file_upload, "read_owned_bytes", read)


@pytest.mark.parametrize(("name", "content"), [
    ("a.xlsx", _xlsx()),
    ("a.xls", (Path(__file__).parent / "fixtures" / "s10_sample.xls").read_bytes()),
    ("a.csv", b"name,qty\napple,2\n"),
    ("a.docx", _docx()),
    ("a.pdf", _pdf()),
    ("a.md", b"# heading\nhello"),
], ids=["xlsx", "xls", "csv", "docx", "pdf", "md"])
def test_all_document_formats(monkeypatch: pytest.MonkeyPatch, name: str, content: bytes) -> None:
    _owned(monkeypatch, name, content)
    result = file_tools.read_file(7, 12)
    assert result["success"] is True
    assert result["fileName"] == name
    assert result["format"] == name.rsplit(".", 1)[1]
    if name.endswith(("xlsx", "xls", "csv")):
        assert result["headers"] == ["name", "qty"]
        assert result["totalRows"] == 1
        assert result["data"]
        assert result["truncated"] is False
    else:
        assert "content" in result
        assert "totalChars" in result
        if name.endswith("docx"):
            assert "item\t3" in str(result["content"])


def test_exact_table_count_and_preview(monkeypatch: pytest.MonkeyPatch) -> None:
    _owned(monkeypatch, "large.xlsx", _xlsx(207))
    result = file_tools.read_file(7, 12)
    assert result["totalRows"] == 207
    assert len(result["data"]) == 200
    assert result["truncated"] is True
    assert "207" in str(result["message"])


def test_csv_header_and_text_limits(monkeypatch: pytest.MonkeyPatch) -> None:
    _owned(monkeypatch, "duplicate.csv", b"name,name,\na,b,c\n")
    assert file_tools.read_file(7, 12)["headers"] == ["name", "name_2", "column_2"]
    _owned(monkeypatch, "long.md", b"x" * 30_001)
    result = file_tools.read_file(7, 12)
    assert len(result["content"]) == 30_000
    assert result["totalChars"] == 30_001
    assert result["truncated"] is True


def test_pdf_page_limit(monkeypatch: pytest.MonkeyPatch) -> None:
    _owned(monkeypatch, "many.pdf", _pdf(51))
    result = file_tools.read_file(7, 12)
    assert result["totalPages"] == 51
    assert result["truncated"] is True


@pytest.mark.parametrize(("name", "content"), [
    ("bad.xlsx", b"not-a-zip"),
    ("bad.xls", b"bad"),
    ("bad.csv", b"\xff"),
    ("bad.docx", b"not-a-zip"),
    ("bad.pdf", b"not-a-pdf"),
    ("bad.md", b"\xff"),
], ids=["xlsx", "xls", "csv", "docx", "pdf", "md"])
def test_corrupt_documents_fail(monkeypatch: pytest.MonkeyPatch, name: str, content: bytes) -> None:
    _owned(monkeypatch, name, content)
    with pytest.raises(ValueError):
        file_tools.read_file(7, 12)


def test_empty_oversize_and_ownership_fail(monkeypatch: pytest.MonkeyPatch) -> None:
    _owned(monkeypatch, "a.csv", b"")
    with pytest.raises(ValueError, match="empty"):
        file_tools.read_file(7, 12)
    _owned(monkeypatch, "a.csv", b"a" * (file_tools.MAX_FILE_BYTES + 1))
    with pytest.raises(BusinessError):
        file_tools.read_file(7, 12)
    _owned(monkeypatch, "a.csv", b"x\ny")
    with pytest.raises(BusinessError):
        file_tools.read_file(8, 12)


_IMAGES = {
    "png": b"\x89PNG\r\n\x1a\nxxIEND",
    "jpg": b"\xff\xd8payload\xff\xd9",
    "jpeg": b"\xff\xd8payload\xff\xd9",
    "webp": b"RIFF\x04\x00\x00\x00WEBP",
    "bmp": b"BM\x0e\x00\x00\x00" + b"\x00" * 8,
}


@pytest.mark.parametrize("extension", list(_IMAGES))
def test_image_formats_and_truncation(monkeypatch: pytest.MonkeyPatch, extension: str) -> None:
    _owned(monkeypatch, f"image.{extension}", _IMAGES[extension])
    monkeypatch.setattr(file_tools, "get_settings", lambda: SimpleNamespace(ocr_enabled=True))
    monkeypatch.setattr(file_tools, "recognize", lambda content: "字" * 30_001)
    result = file_tools.read_image(7, 12)
    assert result["success"] is True
    assert result["fileName"] == f"image.{extension}"
    assert len(result["text"]) == 30_000
    assert result["truncated"] is True


def test_image_disabled_corrupt_size_and_owner(monkeypatch: pytest.MonkeyPatch) -> None:
    _owned(monkeypatch, "image.png", _IMAGES["png"])
    monkeypatch.setattr(file_tools, "get_settings", lambda: SimpleNamespace(ocr_enabled=False))
    with pytest.raises(RuntimeError, match="disabled"):
        file_tools.read_image(7, 12)
    with pytest.raises(BusinessError):
        file_tools.read_image(8, 12)
    _owned(monkeypatch, "image.png", b"corrupt")
    with pytest.raises(ValueError):
        file_tools.read_image(7, 12)
    _owned(monkeypatch, "image.png", b"x" * (file_tools.MAX_IMAGE_BYTES + 1))
    with pytest.raises(BusinessError):
        file_tools.read_image(7, 12)


def test_aliyun_protocol_and_missing_credentials(monkeypatch: pytest.MonkeyPatch) -> None:
    class FakeClient:
        def recognize_advanced(self, request: object) -> object:
            assert request.body.read() == b"image"  # type: ignore[attr-defined]
            return SimpleNamespace(body=SimpleNamespace(data=json.dumps({"content": "识别文本"})))

    monkeypatch.setattr(ocr, "get_settings", lambda: SimpleNamespace(ocr_enabled=True,
                        ocr_access_key_id="", ocr_access_key_secret="", oss_access_key_id="",
                        oss_access_key_secret="", ocr_endpoint="endpoint"))
    assert ocr.recognize(b"image", client_factory=FakeClient) == "识别文本"
    with pytest.raises(RuntimeError, match="credentials"):
        ocr.recognize(b"image")
